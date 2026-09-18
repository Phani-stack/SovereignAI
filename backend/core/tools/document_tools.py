from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as PPTRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from langchain_core.tools import tool


import re
from datetime import datetime

OUTPUT_DIR = Path("./data/outputs")

GENERIC_NAMES = {
    "output", "document", "report", "summary", "file", "untitled", "doc",
    "data", "presentation", "sheet", "export", "test", "demo", "sample",
    "generated_presentation", "generated_document", "generated_file",
    "powerpoint_presentation", "presentation_1"
}

def clean_structural_text(text: str) -> str:
    if not text or not text.strip():
        return ""

    t = text.strip()
    t = re.sub(r'^(slide|chapter|section|page|part|unit)\s*\d*[\s:-]*', '', t, flags=re.IGNORECASE).strip()
    t = re.sub(r'^(executive\s+title\s+overview|title\s+overview|executive\s+title)[\s:-]*', '', t, flags=re.IGNORECASE).strip()

    words = [w for w in re.sub(r'[^\w\s-]', '', t).split()]
    if words:
        return " ".join(words)
    return ""


def derive_relevant_filename(filename: str, title: str = "", content: str = "", default_ext: str = ".docx") -> str:
    p = Path(filename)
    stem = p.stem.strip().lower()
    ext = p.suffix.lower() if p.suffix else default_ext

    is_generic = (not stem or stem in GENERIC_NAMES or stem.startswith("generated_") or len(stem) < 3 or "slide_" in stem)

    if is_generic:
        candidate_phrase = ""
        # 1. Try title first
        if title and title.strip():
            candidate_phrase = clean_structural_text(title)

        # 2. Try content lines if title was structural or empty
        if not candidate_phrase and content and content.strip():
            for line in content.splitlines():
                l = line.strip()
                if not l or l.startswith('#') or l.startswith('---'):
                    continue
                phrase = clean_structural_text(l)
                if phrase:
                    candidate_phrase = phrase
                    break

        if candidate_phrase:
            words = candidate_phrase.split()[:5]
            if words:
                stem = "_".join(words).lower()[:50]

        if not stem or stem in GENERIC_NAMES or stem.startswith("generated_") or "slide_" in stem:
            stem = f"sovereign_ai_{default_ext.lstrip('.')}"

    stem = re.sub(r'[^\w-]', '_', stem).strip('_')
    if not ext.startswith('.'):
        ext = f".{ext}"

    return f"{stem}{ext}"


def prepare_output(filename: str, title: str = "", content: str = "", default_ext: str = ".docx") -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = derive_relevant_filename(filename, title, content, default_ext)

    target_path = OUTPUT_DIR / filename
    if not target_path.exists():
        return target_path

    p = Path(filename)
    stem = p.stem
    ext = p.suffix
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    new_filename = f"{stem}_{timestamp}{ext}"
    target_path = OUTPUT_DIR / new_filename

    counter = 1
    while target_path.exists():
        new_filename = f"{stem}_{timestamp}_{counter}{ext}"
        target_path = OUTPUT_DIR / new_filename
        counter += 1

    return target_path


@tool
def create_docx(filename: str, title: str, content: str) -> str:
    """Create a styled, professional DOCX document with a clear, subject-descriptive filename."""
    try:
        path = prepare_output(filename, title=title, content=content, default_ext=".docx")
        doc = Document()


        # Set 1 inch margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Title Block
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run(title)
        run_title.font.name = 'Calibri'
        run_title.font.size = Pt(24)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)  # #0F172A

        # Accent Line
        accent_p = doc.add_paragraph()
        accent_p.paragraph_format.space_after = Pt(18)
        run_acc = accent_p.add_run("―" * 40)
        run_acc.font.name = 'Calibri'
        run_acc.font.size = Pt(10)
        run_acc.font.color.rgb = RGBColor(59, 130, 246)  # #3B82F6

        # Parse Markdown-style content
        lines = content.splitlines()
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            if stripped.startswith("### "):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(stripped[4:].strip())
                run.font.name = 'Calibri'
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.color.rgb = RGBColor(51, 65, 85)

            elif stripped.startswith("## "):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(stripped[3:].strip())
                run.font.name = 'Calibri'
                run.font.size = Pt(15)
                run.font.bold = True
                run.font.color.rgb = RGBColor(29, 78, 216)  # #1D4ED8

            elif stripped.startswith("# "):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(stripped[2:].strip())
                run.font.name = 'Calibri'
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 58, 138)  # #1E3A8A

            elif stripped.startswith("- ") or stripped.startswith("* "):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(stripped[2:].strip())
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(51, 65, 85)

            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(stripped)
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(51, 65, 85)

        doc.save(path)
        return f"Document created successfully: {path}"

    except Exception as e:
        return f"DOCX creation failed: {e}"


@tool
def create_pdf(filename: str, title: str, content: str) -> str:
    """Create a styled, professional PDF document with a clear, subject-descriptive filename."""
    try:
        path = prepare_output(filename, title=title, content=content, default_ext=".pdf")

        doc = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            leftMargin=40,
            rightMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=8
        )

        h1_style = ParagraphStyle(
            'DocH1',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#1E3A8A'),
            spaceBefore=14,
            spaceAfter=6
        )

        h2_style = ParagraphStyle(
            'DocH2',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=15,
            textColor=colors.HexColor('#2563EB'),
            spaceBefore=10,
            spaceAfter=4
        )

        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155'),
            spaceAfter=6
        )

        bullet_style = ParagraphStyle(
            'DocBullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155'),
            leftIndent=15,
            spaceAfter=4
        )

        story = []
        story.append(Paragraph(title, title_style))
        story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#3B82F6'), spaceBefore=0, spaceAfter=14))

        lines = content.splitlines()
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            if stripped.startswith("### ") or stripped.startswith("## "):
                story.append(Paragraph(stripped.lstrip("#").strip(), h2_style))
            elif stripped.startswith("# "):
                story.append(Paragraph(stripped[2:].strip(), h1_style))
            elif stripped.startswith("- ") or stripped.startswith("* "):
                story.append(Paragraph(f"• {stripped[2:].strip()}", bullet_style))
            else:
                story.append(Paragraph(stripped, body_style))

        doc.build(story)
        return f"PDF created successfully: {path}"

    except Exception as e:
        return f"PDF creation failed: {e}"


@tool
def create_pptx(filename: str, title: str, content: str) -> str:
    """Create a styled, modern executive PowerPoint presentation with a clear, subject-descriptive filename."""
    try:
        path = prepare_output(filename, title=title, content=content, default_ext=".pptx")

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        dark_navy = PPTRGBColor(15, 23, 42)     # #0F172A
        accent_blue = PPTRGBColor(59, 130, 246)  # #3B82F6
        card_bg = PPTRGBColor(30, 41, 59)       # #1E293B
        text_white = PPTRGBColor(248, 250, 252)  # #F8FAFC
        text_muted = PPTRGBColor(148, 163, 184) # #94A3B8

        blank_layout = prs.slide_layouts[6]

        # Slide 1: Title Slide
        title_slide = prs.slides.add_slide(blank_layout)
        background = title_slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = dark_navy

        # Card container on Title Slide
        title_box = title_slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.333), Inches(3.5))
        tf = title_box.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(36)
        p_title.font.bold = True
        p_title.font.color.rgb = text_white
        p_title.alignment = PP_ALIGN.CENTER

        p_sub = tf.add_paragraph()
        p_sub.text = "SovereignAI Executive Presentation"
        p_sub.font.size = Pt(18)
        p_sub.font.color.rgb = accent_blue
        p_sub.alignment = PP_ALIGN.CENTER
        p_sub.space_before = Pt(12)

        # Parse content into structured slides
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        slides_data = []
        current_slide_title = None
        current_bullets = []

        for line in lines:
            if line.startswith("# ") or line.startswith("## ") or line.startswith("### "):
                if current_bullets or current_slide_title:
                    s_name = current_slide_title if current_slide_title else title
                    slides_data.append((s_name, current_bullets if current_bullets else ["Overview of presentation topic."]))
                    current_bullets = []
                current_slide_title = line.lstrip("#").strip()
            else:
                clean_bullet = line.lstrip("-*•1234567890. ").strip()
                if clean_bullet:
                    current_bullets.append(clean_bullet)

        if current_bullets or current_slide_title:
            s_name = current_slide_title if current_slide_title else title
            slides_data.append((s_name, current_bullets if current_bullets else ["Overview of presentation topic."]))

        # Auto-chunking: If a slide has > 4 bullets, split into multiple slides
        expanded_slides = []
        for s_t, s_b in slides_data:
            if len(s_b) <= 4:
                expanded_slides.append((s_t, s_b))
            else:
                # Chunk into groups of 3-4 bullets
                chunk_size = 4
                for c_idx in range(0, len(s_b), chunk_size):
                    chunk_title = f"{s_t} (Part {c_idx // chunk_size + 1})" if len(s_b) > 4 else s_t
                    expanded_slides.append((chunk_title, s_b[c_idx:c_idx + chunk_size]))

        # Guarantee at least 5 slides if request asks for a presentation
        default_sections = [
            ("Executive Summary & Overview", ["High-level context and scope", "Key objectives and target outcomes", "Strategic alignment and goals"]),
            ("Background & Key Concepts", ["Core background information", "Industry context and motivation", "Primary challenges addressed"]),
            ("Detailed Analysis & Features", ["Technical breakdown and core capabilities", "Performance highlights and architecture", "Operational advantages"]),
            ("Strategic Impact & Benefits", ["Measurable results and efficiency gains", "Security, privacy, and sovereignty", "Long-term scalability and value"]),
            ("Summary & Next Steps", ["Summary of core findings", "Actionable recommendations", "Phased implementation roadmap"])
        ]

        # If fewer than 5 slides were generated, pad with structured content from bullets or default sections
        if len(expanded_slides) < 5:
            # Flatten all bullets available
            all_bullets = [b for _, bullets in expanded_slides for b in bullets]
            if len(all_bullets) >= 5:
                # Re-distribute all bullets evenly across 5 slides
                expanded_slides = []
                bullets_per_slide = max(1, len(all_bullets) // 5)
                for i in range(5):
                    sec_title, sec_bullets = default_sections[i]
                    start = i * bullets_per_slide
                    end = (i + 1) * bullets_per_slide if i < 4 else len(all_bullets)
                    slide_b = all_bullets[start:end] if start < len(all_bullets) else sec_bullets
                    expanded_slides.append((f"Slide {i+1}: {sec_title}", slide_b))
            else:
                # Fill remaining slots using default_sections
                while len(expanded_slides) < 5:
                    idx = len(expanded_slides)
                    sec_title, sec_bullets = default_sections[idx % len(default_sections)]
                    expanded_slides.append((f"Slide {idx+1}: {sec_title}", sec_bullets))

        for slide_idx, (s_title, s_bullets) in enumerate(expanded_slides, start=1):
            slide = prs.slides.add_slide(blank_layout)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = dark_navy

            # Slide Title
            header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.0))
            htf = header_box.text_frame
            htf.word_wrap = True
            hp = htf.paragraphs[0]
            hp.text = s_title
            hp.font.size = Pt(24)
            hp.font.bold = True
            hp.font.color.rgb = text_white

            # Accent line shape under header
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.04))
            accent.fill.solid()
            accent.fill.fore_color.rgb = accent_blue
            accent.line.fill.background()

            # Content Card Container
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.1))
            card.fill.solid()
            card.fill.fore_color.rgb = card_bg
            card.line.color.rgb = accent_blue

            # Text inside Card
            content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.933), Inches(4.5))
            ctf = content_box.text_frame
            ctf.word_wrap = True

            for b_idx, bullet in enumerate(s_bullets[:5]):  # max 5 bullet points per slide
                cp = ctf.paragraphs[0] if b_idx == 0 else ctf.add_paragraph()
                cp.text = f"•  {bullet}"
                cp.font.size = Pt(16)
                cp.font.color.rgb = text_white
                cp.space_after = Pt(12)

        prs.save(path)
        return f"PPTX created successfully: {path}"

    except Exception as e:
        return f"PPTX creation failed: {e}"


@tool
def create_xlsx(filename: str, title: str, content: str) -> str:
    """Create a styled, professional Excel XLSX workbook with a clear, subject-descriptive filename."""
    try:
        path = prepare_output(filename, title=title, content=content, default_ext=".xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary Report"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_body = Font(name="Calibri", size=11, color="1E293B")

        fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_alt_row = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        # Title Row across columns A-D
        ws.merge_cells("A1:D1")
        title_cell = ws["A1"]
        title_cell.value = title
        title_cell.font = font_title
        title_cell.fill = fill_title
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 36

        # Sub-header
        ws["A3"] = "Item / Description"
        ws["A3"].font = font_header
        ws["A3"].fill = fill_header
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws["B3"] = "Value / Details"
        ws["B3"].font = font_header
        ws["B3"].fill = fill_header
        ws["B3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[3].height = 24

        lines = [line.strip() for line in content.splitlines() if line.strip()]
        current_row = 4

        for idx, line in enumerate(lines):
            col1 = f"Row {idx + 1}"
            col2 = line

            if ":" in line and not line.startswith("http"):
                parts = line.split(":", 1)
                col1 = parts[0].strip().lstrip("-#* ")
                col2 = parts[1].strip()

            c1 = ws.cell(row=current_row, column=1, value=col1)
            c2 = ws.cell(row=current_row, column=2, value=col2)

            c1.font = font_body
            c2.font = font_body
            c1.border = thin_border
            c2.border = thin_border

            if current_row % 2 == 0:
                c1.fill = fill_alt_row
                c2.fill = fill_alt_row

            ws.row_dimensions[current_row].height = 20
            current_row += 1

        # Adjust column widths
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 75

        wb.save(path)
        return f"XLSX created successfully: {path}"

    except Exception as e:
        return f"XLSX creation failed: {e}"
